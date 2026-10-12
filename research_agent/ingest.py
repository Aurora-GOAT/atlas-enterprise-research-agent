"""Extract searchable plain text from uploaded files.

Design note: text is indexed locally and never sent to any API in offline mode.
"""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
import csv
import io
import re
from uuid import uuid4
from .models import Document

SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".csv"}
MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 120_000


def extract_text(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported format {suffix or '(none)'}; expected {', '.join(sorted(SUPPORTED))}")
    if not payload:
        raise ValueError("File is empty")
    if len(payload) > MAX_BYTES:
        raise ValueError("File exceeds 8 MiB limit")
    if suffix in {".md", ".txt"}:
        text = payload.decode("utf-8-sig", errors="replace")
    elif suffix == ".pdf":
        import fitz
        try:
            with fitz.open(stream=payload, filetype="pdf") as doc:
                if len(doc) > 100:
                    raise ValueError("PDF exceeds 100 pages")
                text = "\n".join(page.get_text("text") for page in doc)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Unable to parse PDF: {exc}") from exc
    elif suffix == ".docx":
        from docx import Document as WordDocument
        try:
            d = WordDocument(BytesIO(payload))
            text = "\n".join(p.text for p in d.paragraphs)
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows)
        except Exception as exc:
            raise ValueError(f"Unable to parse DOCX: {exc}") from exc
    else:
        decoded = payload.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        text = "\n".join(" | ".join(row[:24]) for _, row in zip(range(5000), reader))
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 15:
        raise ValueError("Document contains insufficient extractable text (scanned PDF may need OCR)")
    return text[:MAX_TEXT_CHARS]


def make_document(filename: str, payload: bytes, source_type: str = "upload") -> Document:
    return Document(
        document_id=f"DOC-{uuid4().hex[:10]}",
        title=Path(filename).stem[:100],
        filename=Path(filename).name[:150],
        text=extract_text(filename, payload),
        source_type=source_type,
    )
